"""Alertas de empuje: SKU que van bajo forecast en la semana en curso.

Uso:
    python seguimiento/empuje.py data/S44/foto_2026-10-26.csv [--min-t 1] [--factor 0.5] [--quiebre 0.3]

Lee la foto del día (Cadena, Grupo, Subcategoria, SKU, Producto, FCST, Solicitado, Venta Real, Quebrados)
que deja avance_semana.py y arma reportes/Empuje_S<sem>_<fecha>.xlsx.

Regla (por cadena x SKU con FCST >= min-t):
- "Sin pedidos": no tiene solicitado.
- "Bajo ritmo": su avance (solicitado / FCST) es menor que factor x el avance de su cadena
  y le faltan al menos min-t toneladas.
- Si lo quebrado supera el umbral de quiebre sobre lo solicitado, va a "Quiebre": el problema
  es stock, no demanda, así que no se empuja.
"""
import argparse
import os
import re

import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("foto")
    ap.add_argument("--min-t", type=float, default=1.0)
    ap.add_argument("--factor", type=float, default=0.5)
    ap.add_argument("--quiebre", type=float, default=0.3)
    a = ap.parse_args()

    m = re.search(r"S(\d+)[/\\]foto_(\d{4}-\d{2}-\d{2})", a.foto)
    sem, fecha = m.group(1), m.group(2)
    d = pd.read_csv(a.foto)
    cad = d.groupby("Cadena")[["FCST", "Solicitado", "Venta Real", "Quebrados"]].sum()
    cad["Avance cadena"] = cad["Solicitado"] / cad["FCST"].where(cad["FCST"] > 0)
    d["Avance"] = d["Solicitado"] / d["FCST"].where(d["FCST"] > 0)
    d["Avance cadena"] = d["Cadena"].map(cad["Avance cadena"])
    d["Falta (t)"] = (d["FCST"] - d["Solicitado"]).clip(lower=0)
    d["Quiebre %"] = d["Quebrados"] / d["Solicitado"].where(d["Solicitado"] > 0)

    base = d[(d["FCST"] >= a.min_t)].copy()
    sin = base["Solicitado"] <= 0
    bajo = (base["Avance"] < a.factor * base["Avance cadena"]) & (base["Falta (t)"] >= a.min_t)
    queb = base["Quiebre %"] >= a.quiebre
    base["Alerta"] = None
    base.loc[bajo, "Alerta"] = "Bajo ritmo"
    base.loc[sin, "Alerta"] = "Sin pedidos"
    base.loc[queb & base["Alerta"].notna(), "Alerta"] = "Quiebre (no empujar)"
    al = base[base["Alerta"].notna()].copy()

    al["Qué hacer"] = al["Alerta"].map({
        "Sin pedidos": "Empujar: pedir al KAM que gestione la orden",
        "Bajo ritmo": "Empujar: la cadena va más adelantada que este SKU",
        "Quiebre (no empujar)": "Revisar stock: lo pedido no se está despachando",
    })
    cols = ["Cadena", "Grupo", "Subcategoria", "SKU", "Producto", "FCST", "Solicitado", "Avance", "Avance cadena",
            "Falta (t)", "Venta Real", "Quebrados", "Quiebre %", "Alerta", "Qué hacer"]
    orden = {"Sin pedidos": 0, "Bajo ritmo": 1, "Quiebre (no empujar)": 2}
    al = al.assign(_o=al["Alerta"].map(orden)).sort_values(["Cadena", "_o", "Falta (t)"], ascending=[True, True, False])[cols]

    res = cad.copy()
    for k in ["Sin pedidos", "Bajo ritmo", "Quiebre (no empujar)"]:
        res["SKU " + k.lower()] = al[al["Alerta"] == k].groupby("Cadena").size().reindex(res.index).fillna(0).astype(int)
    res["Falta en SKU a empujar (t)"] = al[al["Alerta"] != "Quiebre (no empujar)"].groupby("Cadena")["Falta (t)"].sum().reindex(res.index).fillna(0)
    res["Falta total (t)"] = (res["FCST"] - res["Solicitado"]).clip(lower=0)
    res = res.sort_values("FCST", ascending=False).reset_index()

    out = os.path.join(ROOT, "reportes", f"Empuje_S{sem}_{fecha}.xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        res.to_excel(w, sheet_name="Resumen por cadena", index=False)
        al[al["Alerta"] != "Quiebre (no empujar)"].to_excel(w, sheet_name="SKU a empujar", index=False)
        al[al["Alerta"] == "Quiebre (no empujar)"].to_excel(w, sheet_name="Quiebre (no empujar)", index=False)
        d.sort_values(["Cadena", "FCST"], ascending=[True, False])[cols[:13]].to_excel(w, sheet_name="Detalle completo", index=False)
        pd.DataFrame({"Parámetro": ["Semana", "Corte", "FCST mínimo", "Bajo ritmo", "Quiebre"],
                      "Valor": [f"S{sem}", fecha, f"{a.min_t} t",
                                f"avance del SKU < {a.factor:.0%} del avance de su cadena y faltan ≥ {a.min_t} t",
                                f"quebrado ≥ {a.quiebre:.0%} de lo solicitado"]}).to_excel(w, sheet_name="Parámetros", index=False)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    pct = {"Avance", "Avance cadena", "Quiebre %"}
    for ws in wb.worksheets:
        heads = [c.value for c in ws[1]]
        for i, h in enumerate(heads, 1):
            letter = ws.cell(1, i).column_letter
            ws.column_dimensions[letter].width = min(46, max(10, len(str(h)) + 2, *(len(str(c.value or "")) + 1 for c in ws[letter][1:200])))
            for c in ws[letter][1:]:
                if isinstance(c.value, (int, float)) and h != "SKU":
                    c.number_format = "0%" if h in pct else "#,##0.0"
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    print(res.round(2).to_string())
    print("Excel:", out)


if __name__ == "__main__":
    main()
