"""FCST vs Solicitado por cadena y subcategoría, en Excel.

Uso:
    python seguimiento/fcst_vs_solicitado.py data/S44/foto_2026-10-26.csv

Arma reportes/FCST_vs_Solicitado_S<sem>_<fecha>.xlsx con:
- Por cadena, Por subcategoría, Cadena x subcategoría (tabla larga con filtros),
- Matriz avance (subcategoría en filas, cadena en columnas, con color),
- SKU sobre FCST y Detalle SKU.
"""
import os
import re
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MED = ["FCST", "Solicitado", "Venta Real", "Quebrados"]


def tabla(g):
    g = g.copy()
    g["Solic. − FCST (t)"] = g["Solicitado"] - g["FCST"]
    g["Avance"] = g["Solicitado"] / g["FCST"].where(g["FCST"] > 0)
    g["Falta (t)"] = (g["FCST"] - g["Solicitado"]).clip(lower=0)
    g["Quiebre %"] = g["Quebrados"] / g["Solicitado"].where(g["Solicitado"] > 0)
    return g.rename(columns={"FCST": "FCST (t)", "Solicitado": "Solicitado (t)", "Venta Real": "Venta Real (t)", "Quebrados": "Quebrado (t)"})


def con_total(t, etiqueta_col):
    tot = t.sum(numeric_only=True)
    fila = {etiqueta_col: "TOTAL", **tot.to_dict()}
    fila["Avance"] = tot["Solicitado (t)"] / tot["FCST (t)"] if tot["FCST (t)"] else None
    fila["Quiebre %"] = tot["Quebrado (t)"] / tot["Solicitado (t)"] if tot["Solicitado (t)"] else None
    return pd.concat([t, pd.DataFrame([fila])], ignore_index=True)


def main():
    foto = sys.argv[1]
    m = re.search(r"S(\d+)[/\\]foto_(\d{4}-\d{2}-\d{2})", foto)
    sem, fecha = m.group(1), m.group(2)
    d = pd.read_csv(foto)
    d["Cadena"] = d["Cadena"].str.title().str.replace(" Y ", " y ")
    d["Subcategoria"] = d["Subcategoria"].fillna("SIN SUBCATEGORÍA")

    cad = con_total(tabla(d.groupby("Cadena", as_index=False)[MED].sum()).sort_values("FCST (t)", ascending=False), "Cadena")
    sub = con_total(tabla(d.groupby("Subcategoria", as_index=False)[MED].sum()).sort_values("FCST (t)", ascending=False), "Subcategoria")
    cxs = tabla(d.groupby(["Cadena", "Subcategoria"], as_index=False)[MED].sum())
    cxs = cxs[(cxs["FCST (t)"] > 0) | (cxs["Solicitado (t)"] > 0)].sort_values(["Cadena", "FCST (t)"], ascending=[True, False])

    orden_cad = cad["Cadena"][:-1].tolist()
    orden_sub = sub["Subcategoria"][:-1].tolist()
    p = d.groupby(["Subcategoria", "Cadena"])[["FCST", "Solicitado"]].sum()
    mat = (p["Solicitado"] / p["FCST"].where(p["FCST"] > 0)).unstack().reindex(index=orden_sub, columns=orden_cad)
    tot_sub = sub.set_index("Subcategoria")["Avance"]
    mat["TOTAL"] = tot_sub.reindex(mat.index)
    mat.loc["TOTAL"] = list(cad.set_index("Cadena")["Avance"].reindex(orden_cad)) + [cad["Avance"].iloc[-1]]
    fc = p["FCST"].unstack().reindex(index=orden_sub, columns=orden_cad)
    fc["TOTAL"] = fc.sum(axis=1)
    fc.loc["TOTAL"] = fc.sum()

    sku = tabla(d[["Cadena", "Grupo", "Subcategoria", "SKU", "Producto"] + MED])
    sobre = sku[sku["Solic. − FCST (t)"] > 0.05].sort_values("Solic. − FCST (t)", ascending=False)
    sku = sku.sort_values(["Cadena", "Subcategoria", "FCST (t)"], ascending=[True, True, False])

    out = os.path.join(ROOT, "reportes", f"FCST_vs_Solicitado_S{sem}_{fecha}.xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        cad.to_excel(w, sheet_name="Por cadena", index=False)
        sub.to_excel(w, sheet_name="Por subcategoría", index=False)
        cxs.to_excel(w, sheet_name="Cadena x subcategoría", index=False)
        mat.reset_index().rename(columns={"Subcategoria": "Avance (Solic./FCST)"}).to_excel(w, sheet_name="Matriz avance", index=False)
        fc.reset_index().rename(columns={"Subcategoria": "FCST (t)"}).to_excel(w, sheet_name="Matriz FCST (t)", index=False)
        sobre.to_excel(w, sheet_name="SKU sobre FCST", index=False)
        sku.to_excel(w, sheet_name="Detalle SKU", index=False)

    wb = load_workbook(out)
    pct = {"Avance", "Quiebre %"}
    bold = Font(bold=True)
    for ws in wb.worksheets:
        heads = [c.value for c in ws[1]]
        matriz = ws.title.startswith("Matriz")
        for i, h in enumerate(heads, 1):
            col = ws.cell(1, i).column_letter
            ws.cell(1, i).font = bold
            ws.column_dimensions[col].width = min(44, max(9, len(str(h)) + 2, *(len(str(c.value or "")) + 1 for c in ws[col][1:300])))
            for c in ws[col][1:]:
                if isinstance(c.value, (int, float)) and h != "SKU":
                    c.number_format = "0%" if (h in pct or ws.title == "Matriz avance") else "#,##0.0"
        last = ws.cell(ws.max_row, 1)
        if last.value == "TOTAL":
            for c in ws[ws.max_row]:
                c.font = bold
                c.fill = PatternFill("solid", fgColor="E7EDFD")
        ws.freeze_panes = "B2" if matriz else "A2"
        if not matriz and ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        if ws.title == "Matriz avance":
            rng = f"B2:{ws.cell(ws.max_row, ws.max_column).coordinate}"
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                                              mid_type="num", mid_value=0.5, mid_color="FFEB84",
                                                              end_type="num", end_value=1, end_color="63BE7B"))
        if ws.title in ("Por cadena", "Por subcategoría", "Cadena x subcategoría"):
            idx = heads.index("Avance") + 1
            col = ws.cell(1, idx).column_letter
            ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}", ColorScaleRule(
                start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=0.5, mid_color="FFEB84",
                end_type="num", end_value=1, end_color="63BE7B"))
    wb.save(out)
    print(cad.round(2).to_string(index=False))
    print("Excel:", out)


if __name__ == "__main__":
    main()
